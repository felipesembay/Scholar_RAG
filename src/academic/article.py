"""Article writing: a section-by-section draft grounded in the project sources.

Each section is drafted with the same contract as the chat (academic.grounded): the LLM only
writes [n] markers, citations are rendered from the library and invalid markers are dropped.
The draft keeps its citations as links (#src-1-3) to the section's own numbered sources, so the
final reference list is rebuilt from the citations still present after the author edits.

The article lives in projects/<p>/article.json.
"""

import io
import json
import os
import re
import zipfile
from datetime import datetime

from academic.citations import bibtex, citation_page, make_citekey, reference_list
from academic.grounded import build_sources, format_sources_for_prompt, render_citations
from academic.library import load_library, role_of
from core.config import load_parameters
from core.llm import chat_completion, writing_model
from core.reranker import rerank
from core.vectorstore import load_vectorstore

ARTICLE_FILE = "article.json"
BACKUP_DIR = "article_backups"
WRITING_MAX_SOURCES = 10
CONTEXT_CHARS = 6000  # other sections' text sent as context, so the draft stays coherent
CONTINUE_CHARS = 5000  # end of the section's own text the continuation picks up from
ANCHOR_RE = re.compile(r"\]\(#src-([\d-]+)\)")
CITATION_LINK_RE = re.compile(r"\[(\(.*?\))\]\(#src-[\d-]+\)")  # "[s. d.]" may appear inside

# Empirical article layout (model: Mattera & Franses, 2025, International Journal of Forecasting)
TEMPLATE = [
    (
        "resumo",
        "Resumo",
        "Objetivo, dados, método, principais resultados e contribuição, em um único parágrafo.",
        200,
        False,
    ),
    (
        "introducao",
        "Introdução",
        "Contexto e motivação do problema, lacuna na literatura, objetivo do trabalho, contribuições e organização do artigo.",
        700,
        True,
    ),
    (
        "literatura",
        "Revisão da literatura",
        "Principais abordagens existentes, agrupadas por linha de pesquisa, com seus resultados e limitações, posicionando este trabalho.",
        900,
        True,
    ),
    (
        "metodologia",
        "Metodologia",
        "Modelos e métodos utilizados, com as hipóteses e etapas de estimação, e o desenho da avaliação (amostras, horizontes, métricas).",
        900,
        True,
    ),
    (
        "dados",
        "Dados",
        "Fonte, período, frequência e tratamento dos dados, com estatísticas descritivas relevantes.",
        400,
        True,
    ),
    (
        "resultados",
        "Resultados",
        "Resultados empíricos e comparação entre os modelos, com a discussão do que eles significam frente à literatura.",
        900,
        True,
    ),
    (
        "conclusao",
        "Conclusão",
        "Síntese dos achados, contribuições, limitações e trabalhos futuros.",
        400,
        False,
    ),
]

WRITING_PROMPT = """Você é coautor de um artigo científico e vai redigir a seção "{section}".

Artigo: {article_title}
Objetivo da seção: {goal}
{notes}Extensão aproximada: {words} palavras.

Regras:
- Use SOMENTE as fontes numeradas abaixo como evidência. Depois de cada afirmação baseada nelas, indique a(s) fonte(s) entre colchetes, por exemplo [1] ou [2, 3].
- Use apenas números de fontes listadas abaixo. Nunca escreva nomes de autores, anos, páginas ou referências por conta própria: eles são inseridos automaticamente a partir dos números.
- Fontes marcadas como RESULTADOS DO AUTOR são resultados desta pesquisa: apresente-os como resultados do trabalho, não como literatura.
- Fontes marcadas como MATERIAL DE APOIO (aulas, slides, códigos, dados) não são literatura publicada: podem orientar a explicação, mas um artigo científico não se apoia nelas. Prefira a literatura; se só o material de apoio sustentar um ponto, indique-o com [n] e escreva [PENDENTE: buscar referência publicada].
- Se as fontes não sustentarem algo que o objetivo pede, não invente: escreva [PENDENTE: o que falta] no ponto correspondente.
- Escreva em prosa acadêmica impessoal, em parágrafos, no idioma do objetivo da seção. Sem título da seção, sem lista de referências e sem comentários sobre o texto.
{synthesis}{context}
Fontes:
{sources}"""

CONTINUE_NOTE = """
Texto já escrito nesta seção — continue exatamente de onde ele parou, mantendo o tom, a
terminologia e o argumento. Não repita, não resuma e não reescreva o que já está escrito;
escreva apenas a continuação, em novos parágrafos:
{written}

{instruction}"""

SYNTHESIS_NOTE = (
    "- Esta seção sintetiza o próprio artigo: baseie-se nas seções já redigidas, "
    "sem citações. Se elas ainda não existirem, escreva [PENDENTE: redigir as demais seções antes].\n"
)


def article_path(project_dir: str) -> str:
    return os.path.join(project_dir, ARTICLE_FILE)


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def new_article(title: str = "", blank: bool = False) -> dict:
    """Default empirical layout, or no sections at all (the author adds them one by one)."""
    if blank:
        return {"title": title, "template": "em-branco", "sections": [], "updated": _now()}
    return {
        "title": title,
        "template": "empirico-ijf",
        "sections": [
            {
                "id": sid,
                "title": t,
                "goal": goal,
                "notes": "",
                "words": words,
                "use_sources": src,
                "draft": "",
                "sources": [],
            }
            for sid, t, goal, words, src in TEMPLATE
        ],
        "updated": _now(),
    }


def load_article(project_dir: str) -> dict:
    path = article_path(project_dir)
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return new_article()


def save_article(project_dir: str, article: dict) -> dict:
    article["updated"] = _now()
    with open(article_path(project_dir), "w", encoding="utf-8") as f:
        json.dump(article, f, ensure_ascii=False, indent=2)
    return article


def reset_article(project_dir: str, title: str = "", blank: bool = False) -> dict:
    """Starts over; the previous article (if it had any text) is kept in article_backups/."""
    path = article_path(project_dir)
    if os.path.exists(path):
        old = load_article(project_dir)
        if old.get("title") or any(s.get("draft", "").strip() for s in old["sections"]):
            folder = os.path.join(project_dir, BACKUP_DIR)
            os.makedirs(folder, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
            with open(os.path.join(folder, f"article-{stamp}.json"), "w", encoding="utf-8") as f:
                json.dump(old, f, ensure_ascii=False, indent=2)
    return save_article(project_dir, new_article(title, blank))


def update_article(project_dir: str, payload: dict) -> dict:
    """Saves title and sections edited by the author; ids keep each section's sources."""
    article = load_article(project_dir)
    previous = {s["id"]: s for s in article["sections"]}
    sections = []
    for incoming in payload.get("sections", article["sections"]):
        sid = (incoming.get("id") or "").strip() or f"secao-{len(sections) + 1}"
        base = previous.get(sid, {"sources": []})
        sections.append(
            {
                "id": sid,
                "title": str(incoming.get("title", base.get("title", ""))).strip() or "Seção",
                "goal": str(incoming.get("goal", base.get("goal", ""))),
                "notes": str(incoming.get("notes", base.get("notes", ""))),
                "words": int(incoming.get("words") or base.get("words") or 500),
                "use_sources": bool(incoming.get("use_sources", base.get("use_sources", True))),
                "draft": str(incoming.get("draft", base.get("draft", ""))),
                # the screen sends sources: [] when the author replaced the draft (old links gone)
                **(
                    {"sources": []}
                    if incoming.get("sources") == []
                    else {
                        "sources": base.get("sources", []),
                        **({"generated": base["generated"]} if base.get("generated") else {}),
                    }
                ),
            }
        )
    article["title"] = str(payload.get("title", article.get("title", "")))
    article["sections"] = sections
    return save_article(project_dir, article)


def plain_text(draft: str) -> str:
    """Draft without citation links, so it can be sent back to the LLM as context."""
    text = CITATION_LINK_RE.sub("", draft)
    return re.sub(r"[ \t]+([.,;:])", r"\1", text)


def _context(article: dict, section_id: str) -> str:
    others = [
        f"### {s['title']}\n{plain_text(s['draft'])}"
        for s in article["sections"]
        if s["id"] != section_id and s.get("draft", "").strip()
    ]
    if not others:
        return ""
    text = "\n\n".join(others)[:CONTEXT_CHARS]
    return (
        "\nSeções já redigidas (apenas para coerência; não as repita e não cite a partir delas):\n"
        f"{text}\n"
    )


def retrieve(project_dir: str, query: str, k: int = WRITING_MAX_SOURCES) -> list:
    """Hybrid retrieval + threshold + rerank, as in the chat, with more sources per call."""
    model_cfg = load_parameters()["Model"]
    threshold = float(model_cfg["similarity_search"].get("threshold_default") or 1.1)
    candidates = max(int(model_cfg.get("rerank_candidates") or k), 2 * k)
    hits = load_vectorstore(project_dir).similarity_search_with_score(query, k=candidates)
    hits = [(doc, score) for doc, score in hits if score <= threshold]
    return rerank(query, hits, k)


def _find_section(article: dict, section_id: str) -> dict:
    section = next((s for s in article["sections"] if s["id"] == section_id), None)
    if section is None:
        raise KeyError(section_id)
    return section


def _write(
    project_dir: str,
    article: dict,
    section: dict,
    query: str,
    words: int,
    continuation: str = "",
    temperature: float = 0.3,
) -> tuple[str, list[dict], dict]:
    """One LLM call: retrieval, prompt, rendered citations. Sources are numbered after the
    ones the section already has, so existing citation links keep pointing at their sources."""
    library = load_library(project_dir)
    style = library.get("style", "abnt")
    hits = retrieve(project_dir, query) if section.get("use_sources", True) else []
    sources = build_sources(hits, library, style)
    if continuation:
        offset = max((s["n"] for s in section.get("sources", [])), default=0)
        for s in sources:
            s["n"] += offset

    prompt = WRITING_PROMPT.format(
        section=section["title"],
        article_title=article.get("title") or "(sem título definido)",
        goal=section["goal"] or section["title"],
        notes=f"Orientações do autor: {section['notes']}\n" if section["notes"].strip() else "",
        words=words,
        synthesis="" if section.get("use_sources", True) else SYNTHESIS_NOTE,
        context=_context(article, section["id"]) + continuation,
        sources=format_sources_for_prompt(sources) or "(nenhuma fonte para esta seção)",
    )
    model = writing_model()
    raw = chat_completion([{"role": "user", "content": prompt}], temperature, model=model)
    text, used, dropped = render_citations(raw, sources, style)
    public = [{k: v for k, v in s.items() if not k.startswith("_")} for s in sources]
    meta = {"at": _now(), "model": model, "style": style, "used": used, "dropped": dropped}
    return text, public, meta


def generate_section(project_dir: str, section_id: str) -> dict:
    """Drafts the section from scratch (replaces the current draft)."""
    article = load_article(project_dir)
    section = _find_section(article, section_id)
    query = " ".join(
        filter(None, [article.get("title"), section["title"], section["goal"], section["notes"]])
    )
    draft, sources, meta = _write(project_dir, article, section, query, section["words"])
    section.update(draft=draft, sources=sources, generated={**meta, "mode": "draft"})
    save_article(project_dir, article)
    return section


def continue_section(
    project_dir: str, section_id: str, instruction: str = "", words: int = 300
) -> dict:
    """Appends to the section, picking up where the current text (author's or generated) stops."""
    article = load_article(project_dir)
    section = _find_section(article, section_id)
    written = plain_text(section.get("draft", "")).strip()
    if not written:
        return generate_section(project_dir, section_id)

    tail = written[-CONTINUE_CHARS:]
    query = " ".join(filter(None, [instruction, section["goal"], tail[-600:]]))
    continuation = CONTINUE_NOTE.format(
        written=("[...] " if len(written) > len(tail) else "") + tail,
        instruction=f"Próximo passo pedido pelo autor: {instruction}\n"
        if instruction.strip()
        else "",
    )
    text, sources, meta = _write(project_dir, article, section, query, words, continuation)
    section["draft"] = section["draft"].rstrip() + "\n\n" + text.strip()
    section["sources"] = section.get("sources", []) + sources
    section["generated"] = {**meta, "mode": "continue"}
    save_article(project_dir, article)
    return section


def cited_entries(article: dict, library: dict) -> list[dict]:
    """Library entries still cited in the drafts (citation links that survived editing)."""
    entries, seen = [], set()
    by_source = library.get("entries", {})
    for section in article["sections"]:
        by_n = {s["n"]: s for s in section.get("sources", [])}
        for anchor in ANCHOR_RE.findall(section.get("draft", "")):
            for n in (int(x) for x in anchor.split("-") if x):
                source = by_n.get(n)
                if not source or source["role"] != "literature" or source["source"] in seen:
                    continue
                # role as of now: a file reclassified as material leaves the references
                if (
                    source["source"] in by_source
                    and role_of(library, source["source"]) == "literature"
                ):
                    seen.add(source["source"])
                    entries.append(by_source[source["source"]])
    return entries


def to_markdown(article: dict, library: dict) -> str:
    style = library.get("style", "abnt")
    parts = [f"# {article['title']}"] if article.get("title") else []
    for section in article["sections"]:
        if section.get("draft", "").strip():
            body = CITATION_LINK_RE.sub(r"\1", section["draft"]).strip()
            parts.append(f"## {section['title']}\n\n{body}")
    entries = cited_entries(article, library)
    if entries:
        parts.append(reference_list(entries, style))
    return "\n\n".join(parts) + "\n"


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _add_runs(paragraph, text: str) -> None:
    for piece in INLINE_RE.split(text):
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece:
            paragraph.add_run(piece)


TABLE_SEPARATOR_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")


def table_rows(block: str) -> list[list[str]] | None:
    """Markdown table -> rows (header first, separator dropped); None if not a table."""
    lines = [line.strip() for line in block.strip().splitlines()]
    if len(lines) < 2 or not all(line.startswith("|") for line in lines):
        return None
    if not TABLE_SEPARATOR_RE.match(lines[1]):
        return None
    rows = []
    for line in [lines[0], *lines[2:]]:
        cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
        rows.append([c.strip().replace("\\|", "|") for c in cells])
    width = max(len(r) for r in rows)
    return [r + [""] * (width - len(r)) for r in rows]


def to_docx(article: dict, library: dict) -> bytes:
    """DOCX from the markdown export: headings, paragraphs, bullets, tables, bold and italic."""
    from docx import Document

    doc = Document()
    for block in re.split(r"\n\s*\n", to_markdown(article, library)):
        block = block.strip()
        if not block:
            continue
        rows = table_rows(block)
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    paragraph = table.cell(r, c).paragraphs[0]
                    _add_runs(paragraph, text)
                    if r == 0:
                        for run in paragraph.runs:
                            run.bold = True
            doc.add_paragraph()
        elif re.match(r"^#{4,6} ", block):
            doc.add_heading(block.lstrip("#").strip(), level=3)
        elif block.startswith("# "):
            doc.add_heading(block[2:].strip(), level=0)
        elif block.startswith("## "):
            doc.add_heading(block[3:].strip(), level=1)
        elif block.startswith("### "):
            doc.add_heading(block[4:].strip(), level=2)
        elif all(line.lstrip().startswith(("- ", "* ")) for line in block.splitlines()):
            for line in block.splitlines():
                _add_runs(doc.add_paragraph(style="List Bullet"), line.lstrip()[2:])
        else:
            _add_runs(doc.add_paragraph(), " ".join(line.strip() for line in block.splitlines()))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------- LaTeX

LATEX_SPECIAL = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
    "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
}  # fmt: skip
LATEX_SPECIAL_RE = re.compile(r"[\\&%$#_{}~^]")
CITATION_SPLIT_RE = re.compile(r"(\[\(.*?\)\]\(#src-[\d-]+\))")
PENDING_SPLIT_RE = re.compile(r"(\[PENDENTE:[^\]]*\])")

LATEX_PREAMBLE = {
    "abnt": r"""\documentclass[12pt,a4paper,article,english,brazil]{abntex2}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage{lmodern}
\usepackage{xcolor}
\usepackage{tabularx}
\usepackage[alf]{abntex2cite}
""",
    "apa": r"""\documentclass[12pt,a4paper]{article}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage{lmodern}
\usepackage{xcolor}
\usepackage{tabularx}
\usepackage[authoryear,round]{natbib}
""",
}


def _latex_escape(text: str) -> str:
    return LATEX_SPECIAL_RE.sub(lambda m: LATEX_SPECIAL[m.group(0)], text)


def _latex_inline(text: str) -> str:
    """Escapes plain text; **bold**, *italic* and [PENDENTE: ...] become LaTeX markup."""
    out = []
    for piece in PENDING_SPLIT_RE.split(text):
        if PENDING_SPLIT_RE.fullmatch(piece):
            out.append(r"\textcolor{red}{\textbf{" + _latex_escape(piece) + "}}")
            continue
        for part in INLINE_RE.split(piece):
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                out.append(r"\textbf{" + _latex_escape(part[2:-2]) + "}")
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                out.append(r"\emph{" + _latex_escape(part[1:-1]) + "}")
            else:
                out.append(_latex_escape(part))
    return "".join(out)


def _latex_cite(anchor: str, by_n: dict, keys: dict, style: str) -> str:
    """#src-1-3 -> \\cite[p.~4]{silva2021} (one call per work, so each keeps its pages)."""
    works: dict[str, list[int]] = {}
    results, material = False, []
    for n in (int(x) for x in anchor.split("-") if x):
        source = by_n.get(n)
        if not source:
            continue
        if source["role"] == "material":
            if source["source"] not in material:
                material.append(source["source"])
            continue
        if source["role"] == "results" or source["source"] not in keys:
            results = results or source["role"] == "results"
            continue
        key, entry = keys[source["source"]]
        page = citation_page(entry, source.get("page"))
        works.setdefault(key, [])
        if page and page not in works[key]:
            works[key].append(page)
    command = r"\cite" if style.startswith("abnt") else r"\citep"
    cites = []
    for key, pages in works.items():
        pages = sorted(pages)
        label = (
            f"p.~{pages[0]}"
            if len(pages) == 1
            else f"p.~{', '.join(map(str, pages))}"
            if pages
            else ""
        )
        cites.append(f"{command}[{label}]{{{key}}}" if label else f"{command}{{{key}}}")
    if material:
        cites.append(_latex_escape(f"(material de apoio: {', '.join(material)})"))
    if results:
        cites.append("(resultados desta pesquisa)")
    return "; ".join(cites)


def _latex_table(rows: list[list[str]]) -> str:
    """tabularx spreads the columns over the text width, so wide Word tables still fit."""
    cols = "|" + "X|" * len(rows[0])
    lines = [
        r"\begin{table}[htbp]",
        r"\centering",
        r"\small",
        rf"\begin{{tabularx}}{{\textwidth}}{{{cols}}}",
        r"\hline",
    ]
    for i, row in enumerate(rows):
        cells = [_latex_inline(c) for c in row]
        if i == 0:
            cells = [rf"\textbf{{{c}}}" if c else c for c in cells]
        lines.append(" & ".join(cells) + r" \\ \hline")
    lines += [r"\end{tabularx}", r"\end{table}"]
    return "\n".join(lines)


def _latex_body(section: dict, keys: dict, style: str) -> str:
    by_n = {s["n"]: s for s in section.get("sources", [])}
    blocks = []
    for block in re.split(r"\n\s*\n", section.get("draft", "").strip()):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        rows = table_rows(block)
        if rows:
            blocks.append(_latex_table(rows))
            continue
        if block.startswith("#"):
            title = block.lstrip("#").strip()
            command = "subsubsection" if block.startswith("####") else "subsection"
            blocks.append(rf"\{command}{{" + _latex_inline(title) + "}")
            continue
        bullet = all(line.lstrip().startswith(("- ", "* ")) for line in lines)
        items = (
            [line.lstrip()[2:] for line in lines]
            if bullet
            else [" ".join(x.strip() for x in lines)]
        )
        rendered = []
        for item in items:
            parts = CITATION_SPLIT_RE.split(item)
            rendered.append(
                "".join(
                    _latex_cite(ANCHOR_RE.search(part).group(1), by_n, keys, style)
                    if CITATION_SPLIT_RE.fullmatch(part)
                    else _latex_inline(part)
                    for part in parts
                )
            )
        if bullet:
            blocks.append(
                "\\begin{itemize}\n"
                + "\n".join(f"  \\item {r}" for r in rendered)
                + "\n\\end{itemize}"
            )
        else:
            blocks.append(rendered[0])
    return "\n\n".join(blocks)


def to_latex(article: dict, library: dict) -> tuple[str, str]:
    """(main.tex, referencias.bib) with real \\cite commands and only the cited works."""
    style = "apa" if library.get("style") == "apa" else "abnt"
    entries = cited_entries(article, library)
    taken = {e["citekey"] for e in entries if e.get("citekey")}
    keyed, keys = [], {}  # keys: library source file -> (citekey, entry)
    for source, entry in library.get("entries", {}).items():
        if not any(entry is e for e in entries):
            continue
        key = entry.get("citekey") or make_citekey(entry, taken)
        taken.add(key)
        keyed.append({**entry, "citekey": key})
        keys[source] = (key, entry)

    title = _latex_inline(article.get("title") or "Título do artigo")
    parts = [
        LATEX_PREAMBLE[style],
        r"\title{" + title + "}",
        r"\author{}",
        r"\date{}",
        "",
        r"\begin{document}",
        r"\maketitle",
        "",
    ]
    for section in article["sections"]:
        if not section.get("draft", "").strip():
            continue
        body = _latex_body(section, keys, style)
        if section["id"] == "resumo":
            env = "resumo" if style == "abnt" else "abstract"
            parts += [rf"\begin{{{env}}}", body, rf"\end{{{env}}}", ""]
        else:
            parts += [r"\section{" + _latex_inline(section["title"]) + "}", "", body, ""]
    if keyed:
        if style == "apa":
            parts.append(r"\bibliographystyle{apalike}")
        parts.append(r"\bibliography{referencias}")
    parts.append(r"\end{document}")
    return "\n".join(parts) + "\n", bibtex(keyed) if keyed else ""


def to_latex_zip(article: dict, library: dict) -> bytes:
    tex, bib = to_latex(article, library)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("main.tex", tex)
        zf.writestr("referencias.bib", bib)
    return buffer.getvalue()
